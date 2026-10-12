# Menyusun berkas Excel daftar akun demo dari JSON keluaran seed-demo-accounts.ts.
#   python scripts/build-demo-accounts-xlsx.py <input.json> <output.xlsx>
# Berkas keluaran memuat KATA SANDI MENTAH: simpan di luar repo dan kirim lewat kanal yang aman.
import json
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

src, dst = sys.argv[1], sys.argv[2]
rows = json.load(open(src, encoding="utf-8"))

ROLE_ORDER = ["Initiator", "Controller", "Reviewer Maintenance", "Reviewer HSSE", "Approver", "Performer", "Admin Regional"]
rows.sort(key=lambda r: (r["terminal"] == "Semua terminal", r["terminal"], ROLE_ORDER.index(r["peran"]) if r["peran"] in ROLE_ORDER else 99))

NAVY = "0F3B59"
font = lambda **k: Font(name="Arial", size=k.pop("size", 10), **k)
thin = Side(style="thin", color="C5CED3")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()

# ------------------------------------------------------------ Petunjuk
ws = wb.active
ws.title = "Petunjuk"
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 30
ws.column_dimensions["C"].width = 95

r = 2
ws.cell(r, 2, "Akun Demo — Portal PSI & MOC Regional Jatimbalinus").font = font(size=16, bold=True, color=NAVY)
r += 1
ws.cell(r, 2, "RAHASIA: berkas ini memuat kata sandi. Kirim hanya lewat kanal aman dan jangan diteruskan.").font = font(size=10, bold=True, color="B91C1C")
r += 2

def block(title, items):
    global r
    c = ws.cell(r, 2, title)
    c.font = font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=NAVY)
    ws.cell(r, 3).fill = PatternFill("solid", fgColor=NAVY)
    r += 1
    for k, v in items:
        a = ws.cell(r, 2, k)
        a.font = font(bold=True)
        a.alignment = Alignment(vertical="top", wrap_text=True)
        b = ws.cell(r, 3, v)
        b.font = font()
        b.alignment = Alignment(vertical="top", wrap_text=True)
        for cc in (a, b):
            cc.border = border
        ws.row_dimensions[r].height = max(18, 15 * (1 + len(v) // 95))
        r += 1
    r += 1

block("Cara masuk", [
    ("Alamat", "https://dashboard-pertamina.vercel.app"),
    ("Masuk dengan", "Username ATAU email pada lembar \"Akun Demo\", lalu kata sandinya. Huruf kecil semua pada username."),
    ("Isi kata sandi", "Salin-tempel dari kolom Kata sandi. Format: tiga blok empat karakter dipisah tanda strip."),
    ("Autentikasi dua faktor", "Tidak aktif pada akun demo. Pengguna boleh mengaktifkannya sendiri di Pengaturan > Keamanan Akun."),
])
block("Isi lembar \"Akun Demo\"", [
    ("Per terminal", "Setiap terminal (29) memiliki 6 akun: Initiator, Controller, Reviewer Maintenance, Reviewer HSSE, Approver, dan Performer."),
    ("Admin Regional", "Satu akun (demo-regional) untuk seluruh terminal."),
    ("Total", f"{len(rows)} akun. Tanda pengenal akun demo: username berawalan \"demo-\" dan email berakhiran @demo.invalid."),
])
block("Hak setiap peran", [
    ("Initiator", "Melihat semua terminal; mengunggah, mengubah data, mengganti berkas, memindahkan dokumen terminalnya sendiri; memindahkan ke Sampah, memulihkan, dan menghapus permanen dokumen terminalnya; mengajukan MOC."),
    ("Controller", "Hak dokumen sama dengan Initiator; mendaftarkan MOC (memberi Nomor MoC). Tidak melihat tombol MOC Baru."),
    ("Reviewer Maintenance / HSSE", "Peran MOC: memberi review pada pengajuan MOC terminalnya. Muncul di pilihan Reviewer saat Initiator mengajukan MOC. Hak dokumen sama dengan Initiator."),
    ("Approver", "Peran MOC (Terminal Manager): memutuskan pengajuan MOC dan menunjuk Performer."),
    ("Performer", "Peran MOC: pelaksana perubahan yang ditunjuk Approver."),
    ("Admin Regional", "Melihat dan mengunggah dokumen semua terminal; Kelola User; Log Aktivitas; Arsip Registry semua terminal. MOC hanya dapat DILIHAT (keputusan rapat 18 September 2026): tidak mengajukan atau menyetujui MOC."),
])
block("Mencoba alur MOC lengkap di satu terminal", [
    ("Langkah", "1) Masuk sebagai Initiator terminal itu, buka MOC > MOC Baru, isi formulir, Simpan Draf. 2) Buka MOC itu, pilih Reviewer dan Approver dari terminal yang sama, lalu Ajukan. 3) Masuk sebagai tiap Reviewer: Beri review (Sesuai). 4) Masuk sebagai Approver: Putuskan > setujui dan tunjuk Performer. 5) Masuk sebagai Controller: Daftarkan (beri Nomor MoC)."),
    ("Catatan", "Reviewer, Approver, dan Performer hanya muncul untuk terminal yang sama dengan MOC-nya."),
])
block("Perhatian", [
    ("Akun sementara", "Akun demo dibuat untuk uji coba dan akan dihapus sebelum sistem dipakai sungguhan."),
    ("Data bersama", "Data yang dibuat dengan akun demo tersimpan di database yang sama dengan data lain. Hati-hati dengan Hapus permanen: semua pengguna terminal dapat melakukannya dan tidak dapat dibatalkan."),
    ("Sampah", "Dokumen yang dipindahkan ke Sampah dapat dipulihkan selama 30 hari."),
])

# ------------------------------------------------------------ Akun Demo
wa = wb.create_sheet("Akun Demo")
headers = ["Terminal", "Peran", "Nama akun", "Username", "Email", "Kata sandi", "Peran MOC", "Fungsi MOC", "Hak utama"]
widths = [26, 22, 38, 34, 44, 18, 14, 20, 70]
for i, (h, w) in enumerate(zip(headers, widths), start=1):
    c = wa.cell(1, i, h)
    c.font = font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(vertical="center")
    c.border = border
    wa.column_dimensions[get_column_letter(i)].width = w
for n, a in enumerate(rows, start=2):
    vals = [a["terminal"], a["peran"], a["nama"], a["username"], a["email"], a["password"], ", ".join(a.get("mocRoles") or []) or "-", a.get("mocFunction") or "-", a["keterangan"]]
    for i, v in enumerate(vals, start=1):
        c = wa.cell(n, i, v)
        c.font = Font(name="Consolas", size=10, bold=True) if i == 6 else font()
        c.border = border
        c.number_format = "@"
    if n % 2 == 1:
        for i in range(1, len(headers) + 1):
            wa.cell(n, i).fill = PatternFill("solid", fgColor="F3F6F8")
wa.freeze_panes = "A2"
wa.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
wa.row_dimensions[1].height = 22

wb.save(dst)
print("tersimpan:", dst, "| akun:", len(rows))
